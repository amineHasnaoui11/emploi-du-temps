# -*- coding: utf-8 -*-
import openpyxl, re, json, unicodedata

DAYS = ['Lundi','Mardi','Mercredi','Jeudi','Vendredi','Samedi','Dimanche']
COLS = ['G','H','I','J','K','L','M']

AR = {'عربية':'Arabe','رياضيات':'Math','إيقاظ علمي':'Éveil scientifique','فيزياء':'Physique'}

def norm_time(t):
    t = t.replace('H','h').strip()
    m = re.match(r'^(\d{1,2})h(\d{2})?h?$', t)
    if not m: return None
    h = int(m.group(1)); mi = int(m.group(2) or 0)
    return h*60+mi

def fmt(mins):
    return f"{mins//60:02d}:{mins%60:02d}"

TIME_RE = re.compile(r'(\d{1,2}[hH]\d{0,2}[hH]?)\s*[-–]\s*(\d{1,2}[hH]\d{0,2}[hH]?)')
TEACH_RE = re.compile(r'\(([^)]*)\)')
QUINZ_RE = re.compile(r'/\s*[Pp]ar\s*quin?z?aine|/\s*[Pp]ar\s*quizaine', re.I)

def parse_cell(txt):
    raw = re.sub(r'\s+',' ',txt).strip()
    if not raw: return None
    rec = {'raw': raw}
    rec['quinzaine'] = bool(QUINZ_RE.search(raw))
    s = QUINZ_RE.sub(' ', raw)
    tm = TIME_RE.search(s)
    if tm:
        rec['start'] = norm_time(tm.group(1)); rec['end'] = norm_time(tm.group(2))
        rec['time_raw'] = tm.group(0)
        s = s[:tm.start()] + ' ' + s[tm.end():]
    else:
        rec['start']=rec['end']=None; rec['time_raw']=None
    te = TEACH_RE.search(s)
    if te:
        rec['teacher'] = te.group(1).strip()
        s = s[:te.start()] + ' | ' + s[te.end():]
    else:
        rec['teacher'] = None
    parts = [p.strip() for p in s.split('|')]
    subj = re.sub(r'\s+',' ',parts[0]).strip()
    note = re.sub(r'\s+',' ',' '.join(parts[1:])).strip() if len(parts)>1 else ''
    rec['subject_raw'] = subj
    rec['subject'] = AR.get(subj, subj)
    rec['note'] = note
    return rec

def parse_sheet(ws):
    # class blocks from column F
    fs = [(c.row, re.sub(r'\s+',' ',str(c.value)).strip()) for c in ws['F'] if c.value and str(c.value).strip()]
    blocks=[]
    for i,(r,name) in enumerate(fs):
        end = fs[i+1][0]-1 if i+1 < len(fs) else ws.max_row
        blocks.append((name, r, end))
    recs=[]
    for name, r0, r1 in blocks:
        for ci,col in enumerate(COLS):
            for r in range(r0, r1+1):
                v = ws[f'{col}{r}'].value
                if v is None or not str(v).strip(): continue
                rec = parse_cell(str(v))
                if rec is None: continue
                rec.update({'classe':name,'day':DAYS[ci],'cell':f'{col}{r}'})
                recs.append(rec)
    hdr = {}
    for ci,col in enumerate(COLS):
        hdr[DAYS[ci]] = re.sub(r'\s+',' ',str(ws[f'{col}2'].value or '')).strip()
    return blocks, hdr, recs

wb = openpyxl.load_workbook('orig.xlsx')
data={}
for ws in wb.worksheets:
    blocks, hdr, recs = parse_sheet(ws)
    data[ws.title]={'blocks':[{'classe':b[0],'r0':b[1],'r1':b[2]} for b in blocks],'header':hdr,'sessions':recs}
    print(f"=== {ws.title!r}: {len(blocks)} classes, {len(recs)} séances")
    print("   header:", hdr)
json.dump(data, open('parsed.json','w',encoding='utf-8'), ensure_ascii=False, indent=1)
