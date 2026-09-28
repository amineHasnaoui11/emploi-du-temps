# -*- coding: utf-8 -*-
import openpyxl, re
wb=openpyxl.load_workbook('orig.xlsx')
DAYS=['Lundi','Mardi','Mercredi','Jeudi','Vendredi','Samedi','Dimanche']
COLS=['G','H','I','J','K','L','M']
for ws in wb.worksheets:
    print("="*100); print("FEUILLE:", repr(ws.title))
    # dernière ligne réellement utilisée
    last=0
    for row in ws.iter_rows(min_row=1,max_row=ws.max_row,min_col=6,max_col=13):
        for c in row:
            if c.value not in (None,'') and str(c.value).strip(): last=max(last,c.row)
    fs=[(c.row, re.sub(r'\s+',' ',str(c.value)).strip()) for c in ws['F'] if c.value and str(c.value).strip()]
    blocks=[]
    for i,(r,n) in enumerate(fs):
        end = fs[i+1][0]-1 if i+1<len(fs) else last
        blocks.append((n,r,end))
    print("  dernière ligne utile:", last)
    # index des fusions
    merges={}
    for m in ws.merged_cells.ranges:
        merges[(m.min_col,m.min_row)]=(m.min_row,m.max_row)
    nshared=0
    for ci,col in enumerate(COLS,7):
        for r in range(1,last+1):
            v=ws.cell(r,ci).value
            if v in (None,'') or not str(v).strip(): continue
            r0,r1 = merges.get((ci,r),(r,r))
            over=[(n,max(0,min(r1,e)-max(r0,s)+1)) for n,s,e in blocks if not (r1<s or r0>e)]
            if len(over)>1:
                nshared+=1
                txt=re.sub(r'\s+',' ',str(v)).strip()
                print(f"  {COLS[ci-7]}{r} (fusion {r0}-{r1}) {DAYS[ci-7]:<9} -> {[f'{n}({k}l)' for n,k in over]}")
                print(f"        {txt}")
    print(f"  => {nshared} cellule(s) débordant sur plusieurs classes")
