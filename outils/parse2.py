# -*- coding: utf-8 -*-
"""Parseur conscient des fusions : une cellule fusionnée qui déborde sur
plusieurs blocs de classes encode une SÉANCE PARTAGÉE entre ces groupes."""
import openpyxl, re, json
DAYS=['Lundi','Mardi','Mercredi','Jeudi','Vendredi','Samedi','Dimanche']
COLS=['G','H','I','J','K','L','M']
AR={'عربية':'Arabe','رياضيات':'Math','إيقاظ علمي':'Éveil scientifique','فيزياء':'Physique'}
TIME_RE=re.compile(r'(\d{1,2}[hH]\d{0,2}[hH]?)\s*[-–]\s*(\d{1,2}[hH]\d{0,2}[hH]?)')
TEACH_RE=re.compile(r'\(([^)]*)\)')
QUINZ_RE=re.compile(r'/\s*[Pp]ar\s*quin?z?aine|/\s*[Pp]ar\s*quizaine',re.I)

def norm_time(t):
    t=t.replace('H','h').strip(); m=re.match(r'^(\d{1,2})h(\d{2})?h?$',t)
    return int(m.group(1))*60+int(m.group(2) or 0) if m else None

def parse_cell(txt):
    raw=re.sub(r'\s+',' ',txt).strip()
    if not raw: return None
    rec={'raw':raw,'quinzaine':bool(QUINZ_RE.search(raw))}
    s=QUINZ_RE.sub(' ',raw)
    tm=TIME_RE.search(s)
    if tm:
        rec['start'],rec['end']=norm_time(tm.group(1)),norm_time(tm.group(2))
        s=s[:tm.start()]+' '+s[tm.end():]
    else: rec['start']=rec['end']=None
    te=TEACH_RE.search(s)
    if te: rec['teacher']=te.group(1).strip(); s=s[:te.start()]+' | '+s[te.end():]
    else: rec['teacher']=None
    parts=[p.strip() for p in s.split('|')]
    subj=re.sub(r'\s+',' ',parts[0]).strip()
    rec['subject_raw']=subj; rec['subject']=AR.get(subj,subj)
    rec['note']=re.sub(r'\s+',' ',' '.join(parts[1:])).strip() if len(parts)>1 else ''
    return rec

wb=openpyxl.load_workbook('orig.xlsx')
data={}
for ws in wb.worksheets:
    last=0
    for row in ws.iter_rows(min_row=1,max_row=ws.max_row,min_col=6,max_col=13):
        for c in row:
            if c.value not in (None,'') and str(c.value).strip(): last=max(last,c.row)
    fs=[(c.row,re.sub(r'\s+',' ',str(c.value)).strip()) for c in ws['F'] if c.value and str(c.value).strip()]
    blocks=[(n,r,(fs[i+1][0]-1 if i+1<len(fs) else last)) for i,(r,n) in enumerate(fs)]
    merges={(m.min_col,m.min_row):(m.min_row,m.max_row) for m in ws.merged_cells.ranges}
    recs=[]
    for ci,col in enumerate(COLS,7):
        for r in range(blocks[0][1],last+1):   # on démarre au 1er bloc de classes (on saute les en-têtes)
            v=ws.cell(r,ci).value
            if v in (None,'') or not str(v).strip(): continue
            rec=parse_cell(str(v))
            if rec is None: continue
            r0,r1=merges.get((ci,r),(r,r))
            over=[(n,min(r1,e)-max(r0,s)+1) for n,s,e in blocks if not (r1<s or r0>e)]
            rec.update({'day':DAYS[ci-7],'cell':f'{col}{r}','merge':[r0,r1],
                        'classes':[n for n,_ in over],'overlaps':{n:k for n,k in over}})
            recs.append(rec)
    hdr={DAYS[i]:re.sub(r'\s+',' ',str(ws[f'{c}2'].value or '')).strip() for i,c in enumerate(COLS)}
    data[ws.title]={'blocks':[{'classe':b[0],'r0':b[1],'r1':b[2]} for b in blocks],'header':hdr,'sessions':recs}
    ns=sum(1 for s in recs if len(s['classes'])>1)
    print(f"{ws.title!r}: {len(recs)} séances dont {ns} partagées | lignes classe-séance = {sum(len(s['classes']) for s in recs)}")
json.dump(data,open('parsed2.json','w',encoding='utf-8'),ensure_ascii=False,indent=1)

# unicité de la clé de rythme
from collections import Counter
for t,d in data.items():
    c=Counter((s['day'],s['start'],s['subject'],s['teacher']) for s in d['sessions'])
    dup=[k for k,v in c.items() if v>1]
    if dup: print("  !! clé non unique:",t,dup)
